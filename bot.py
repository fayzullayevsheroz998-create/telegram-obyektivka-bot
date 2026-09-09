import os, io, sqlite3, zipfile, re
from pathlib import Path
from copy import deepcopy
from dotenv import load_dotenv
from docx import Document
from telegram import Update, ReplyKeyboardMarkup, ReplyKeyboardRemove, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import Application, CommandHandler, MessageHandler, CallbackQueryHandler, ContextTypes, filters

BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / 'template.docx'
OUT = BASE / 'generated'; OUT.mkdir(exist_ok=True)
DB_PATH = BASE / 'bot.db'
load_dotenv(BASE / '.env')
BOT_TOKEN = os.getenv('BOT_TOKEN','').strip()
CARD_NUMBER = os.getenv('CARD_NUMBER','').strip()
PRICE = os.getenv('PRICE','').strip()
ADMIN_IDS = {int(x.strip()) for x in os.getenv('ADMIN_IDS','').split(',') if x.strip().isdigit()}

TEXTS = {
'ru': {
'choose':'Выберите язык интерфейса:', 'welcome':'Здравствуйте! 👋\nЯ помогу автоматически заполнить объективку и отправлю готовый Word-документ.',
'payment':'💳 Для начала работы оплатите услугу переводом на карту.\n\nСумма: {price}\nКарта: {card}\n\nПосле оплаты нажмите «💳 Я оплатил» и отправьте чек/скриншот. После проверки администратором заполнение будет открыто.',
'paid':'💳 Я оплатил','receipt':'Отправьте чек или скриншот оплаты 📷. Можно отправить фото или PDF/файл.','receipt_ok':'✅ Чек получен и отправлен на проверку. Ожидайте подтверждения администратора.','checking':'⏳ Ваш платёж уже находится на проверке.','approved':'✅ Оплата подтверждена!\n\nВведите Ф.И.О. человека:','rejected':'❌ Платёж не подтверждён. Проверьте чек и отправьте его повторно.','disabled':'⏸ Бот временно остановлен для новых заказов. Попробуйте позже.',
'name':'Введите Ф.И.О. человека:','birth_date':'Введите дату рождения (например: 30.11.2006):','birth_place':'Введите место рождения:','nationality':'Введите национальность:','party':'Введите партийность (если нет — «нет»):','education':'Введите образование (например: высшее или среднее):','graduated':'Введите учебное заведение, которое окончили:','specialty':'Введите специальность (если нет — «нет»):','degree':'Введите учёную степень (если нет — «нет»):','title':'Введите учёное звание (если нет — «нет»):','languages':'Какие иностранные языки знает? (если нет — «нет»):','awards':'Есть ли государственные награды? Если нет — «нет»:','deputy':'Является ли депутатом или членом выборного органа? Если нет — «нет»:','work':'Введите трудовую деятельность (если нет — «нет»):','photo':'Отправьте фотографию человека 📷\n\nЕсли фото не нужно, нажмите «⏭ Пропустить».','rel_count':'Сколько близких родственников добавить? Введите число от 0 до 30:','rel_count_error':'Введите целое число от 0 до 30.','relation':'Кем приходится родственник №{i}? Например: Отец, Мать, Брат:','rel_name':'Введите Ф.И.О. родственника:','rel_birth':'Введите дату и место рождения родственника:','rel_work':'Введите место работы и должность:','rel_address':'Введите место жительства:',
'done':'✅ Готово! Документ заполнен.','error':'Произошла ошибка при создании документа: {e}',
'start':'▶️ Start','new':'🔄 Новая объективка','language':'🌐 Сменить язык','menu':'🏠 В главное меню','no':'❌ Нет','skip':'⏭ Пропустить','main':'🏠 Главное меню','lang_ru':'🇷🇺 Русский','lang_uz':'🇺🇿 O‘zbekcha',
'admin_only':'Эта команда доступна только администратору.','admin_panel':'⚙️ Панель администратора\n\nСтатус: {status}\nЗаказов на проверке: {pending}','on':'▶️ Включить','off':'⏹ Выключить','on_status':'🟢 ВКЛЮЧЕН','off_status':'🔴 ОСТАНОВЛЕН','approve':'Подтвердить','reject':'Отклонить','approved_admin':'✅ Оплата подтверждена. Пользователь получил доступ.','rejected_admin':'❌ Оплата отклонена. Пользователь уведомлён.','menu_hint':'Используйте кнопки меню ниже.'},
'uz': {
'choose':'Interfeys tilini tanlang:','welcome':'Assalomu alaykum! 👋\nMen obyektivkani avtomatik to‘ldirib, tayyor Word hujjatini yuboraman.',
'payment':'💳 Ishni boshlash uchun xizmat haqini karta orqali o‘tkazing.\n\nSumma: {price}\nKarta: {card}\n\nTo‘lovdan so‘ng «💳 To‘ladim» tugmasini bosing va chek/skrinshotni yuboring. Administrator tekshirganidan keyin to‘ldirish ochiladi.',
'paid':'💳 To‘ladim','receipt':'To‘lov cheki yoki skrinshotini yuboring 📷. Rasm yoki PDF/fayl yuborish mumkin.','receipt_ok':'✅ Chek qabul qilindi va tekshiruvga yuborildi. Administrator tasdig‘ini kuting.','checking':'⏳ To‘lovingiz allaqachon tekshiruvda.','approved':'✅ To‘lov tasdiqlandi!\n\nShaxsning F.I.Sh. ni kiriting:','rejected':'❌ To‘lov tasdiqlanmadi. Chekni tekshirib, qayta yuboring.','disabled':'⏸ Bot yangi buyurtmalar uchun vaqtincha to‘xtatilgan. Keyinroq urinib ko‘ring.',
'name':'Shaxsning F.I.Sh. ni kiriting:','birth_date':'Tug‘ilgan sanani kiriting (masalan: 30.11.2006):','birth_place':'Tug‘ilgan joyini kiriting:','nationality':'Millatini kiriting:','party':'Partiyaviyligini kiriting (agar yo‘q bo‘lsa — «yo‘q»):','education':'Ma’lumotingizni kiriting (masalan: oliy yoki o‘rta):','graduated':'Tamomlagan ta’lim muassasasini kiriting:','specialty':'Mutaxassisligini kiriting (agar yo‘q bo‘lsa — «yo‘q»):','degree':'Ilmiy darajasini kiriting (agar yo‘q bo‘lsa — «yo‘q»):','title':'Ilmiy unvonini kiriting (agar yo‘q bo‘lsa — «yo‘q»):','languages':'Qaysi chet tillarini biladi? (agar bilmasa — «yo‘q»):','awards':'Davlat mukofotlari bormi? Agar yo‘q bo‘lsa — «yo‘q»:','deputy':'Deputat yoki saylanadigan organ a’zosimi? Agar yo‘q bo‘lsa — «yo‘q»:','work':'Mehnat faoliyatini kiriting (agar yo‘q bo‘lsa — «yo‘q»):','photo':'Shaxsning fotosuratini yuboring 📷\n\nAgar foto kerak bo‘lmasa, «⏭ O‘tkazib yuborish» tugmasini bosing.','rel_count':'Nechta yaqin qarindosh qo‘shilsin? 0 dan 30 gacha raqam kiriting:','rel_count_error':'0 dan 30 gacha bo‘lgan butun sonni kiriting.','relation':'Qarindoshi №{i} kim? Masalan: Otasi, Onasi, Akasi:','rel_name':'Qarindoshining F.I.Sh. ni kiriting:','rel_birth':'Qarindoshining tug‘ilgan sanasi va joyini kiriting:','rel_work':'Ish joyi va lavozimini kiriting:','rel_address':'Yashash manzilini kiriting:',
'done':'✅ Tayyor! Hujjat to‘ldirildi.','error':'Hujjatni yaratishda xatolik yuz berdi: {e}',
'start':'▶️ Start','new':'🔄 Yangi obyektivka','language':'🌐 Tilni almashtirish','menu':'🏠 Bosh menyuga o‘tish','no':'❌ Yo‘q','skip':'⏭ O‘tkazib yuborish','main':'🏠 Bosh menyu','lang_ru':'🇷🇺 Русский','lang_uz':'🇺🇿 O‘zbekcha',
'admin_only':'Bu buyruq faqat administrator uchun.','admin_panel':'⚙️ Administrator paneli\n\nHolat: {status}\nTekshiruvdagi buyurtmalar: {pending}','on':'▶️ Yoqish','off':'⏹ To‘xtatish','on_status':'🟢 YOQILGAN','off_status':'🔴 TO‘XTATILGAN','approve':'Tasdiqlash','reject':'Rad etish','approved_admin':'✅ To‘lov tasdiqlandi. Foydalanuvchiga kirish berildi.','rejected_admin':'❌ To‘lov rad etildi. Foydalanuvchiga xabar berildi.','menu_hint':'Quyidagi menyu tugmalaridan foydalaning.'}
}

STEPS=['name','birth_date','birth_place','nationality','party','education','graduated','specialty','degree','title','languages','awards','deputy','work','photo','rel_count']
OPTIONAL={'party','specialty','degree','title','languages','awards','deputy','work'}
REL_STEPS=['relation','rel_name','rel_birth','rel_work','rel_address']

def normalize(s):
    return (s or '').strip().casefold().replace('’',"'").replace('‘',"'").replace('`',"'").replace('ʻ',"'").replace('ʼ',"'").replace('о\u02bc','о\'').replace('o\u02bc','o\'')

def tr(context,key,**kw): return TEXTS.get(context.user_data.get('lang','ru'),'ru').get(key,key).format(**kw)
def labels(context):
    lang=context.user_data.get('lang','ru'); t=TEXTS[lang]; return {k:normalize(t[k]) for k in ['start','new','language','menu','no','skip','paid']}

def db():
    c=sqlite3.connect(DB_PATH); c.row_factory=sqlite3.Row
    c.execute('CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL)')
    c.execute('CREATE TABLE IF NOT EXISTS orders(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER NOT NULL,username TEXT,lang TEXT,status TEXT NOT NULL,created_at DATETIME DEFAULT CURRENT_TIMESTAMP)')
    c.commit(); return c

def get_setting(k,d=''): c=db(); r=c.execute('SELECT value FROM settings WHERE key=?',(k,)).fetchone(); c.close(); return r['value'] if r else d
def set_setting(k,v): c=db(); c.execute('INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',(k,v)); c.commit(); c.close()
def enabled(): return get_setting('enabled','1')=='1'
def get_order(uid): c=db(); r=c.execute('SELECT * FROM orders WHERE user_id=? ORDER BY id DESC LIMIT 1',(uid,)).fetchone(); c.close(); return r
def create_order(user,lang): c=db(); cur=c.execute('INSERT INTO orders(user_id,username,lang,status) VALUES(?,?,?,?)',(user.id,user.username or '',lang,'pending')); c.commit(); oid=cur.lastrowid; c.close(); return oid
def update_order(oid,status): c=db(); c.execute('UPDATE orders SET status=? WHERE id=?',(status,oid)); c.commit(); c.close()

def menu_kb(lang):
    t=TEXTS[lang]
    return ReplyKeyboardMarkup([[t['start'],t['new']],[t['language']]],resize_keyboard=True)
def form_kb(lang,optional=False,photo=False):
    t=TEXTS[lang]; rows=[]
    if optional: rows.append([t['no'],t['skip']])
    elif photo: rows.append([t['skip']])
    rows.append([t['menu']])
    return ReplyKeyboardMarkup(rows,resize_keyboard=True)
def payment_kb(lang):
    t=TEXTS[lang]; return ReplyKeyboardMarkup([[t['paid']],[t['menu']]],resize_keyboard=True)

def set_para(p,text,bold=False):
    r=p.runs[0] if p.runs else p.add_run(); r.text=str(text); r.bold=bold
    for x in p.runs[1:]: x.text=''
def set_cell(cell,text): cell.text=str(text)

def build_document(data,photo=None):
    doc=Document(TEMPLATE); p=doc.paragraphs
    set_para(p[4],data['name'],True); set_para(p[7],f"{data['birth_date']}                                                    {data['birth_place']}")
    set_para(p[10],f"{data['nationality']}\t\t\t\t\t\t{data['party']}"); set_para(p[13],f"{data['education']}\t{data['graduated']}\t")
    set_para(p[15],data['specialty']); set_para(p[17],f"{data['degree']}\t\t\t\t\t\t{data['title']}")
    set_para(p[20],data['languages']); set_para(p[22],data['awards']); set_para(p[25],data['deputy']); set_para(p[30],data['work'])
    if len(p)>46: set_para(p[46],f"{data['name']}ning yaqin qarindoshlari haqida")
    table=doc.tables[0]
    while len(table.rows)>1: table._tbl.remove(table.rows[-1]._tr)
    template_row=deepcopy(table.rows[0]._tr)
    for rel in data.get('relatives',[]):
        table._tbl.append(deepcopy(template_row)); row=table.rows[-1]
        for c,val in zip(row.cells,[rel['relation'],rel['name'],rel['birth'],rel['work'],rel['address']]): set_cell(c,val)
    out=io.BytesIO(); doc.save(out); raw=out.getvalue()
    if not photo: return raw
    src=OUT/'_photo_tmp.docx'; src.write_bytes(raw); final=io.BytesIO()
    with zipfile.ZipFile(src,'r') as zin, zipfile.ZipFile(final,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            content=zin.read(item.filename)
            if item.filename.startswith('word/media/') and item.filename.lower().endswith(('.png','.jpg','.jpeg')):
                content=photo
            zout.writestr(item,content)
    src.unlink(missing_ok=True); return final.getvalue()

async def show_menu(update,context):
    lang=context.user_data.get('lang','ru'); await update.message.reply_text(TEXTS[lang]['main'],reply_markup=menu_kb(lang))

async def start_cmd(update,context):
    context.user_data.setdefault('lang','ru')
    if not update.message: return
    if update.effective_user.id in ADMIN_IDS:
        await update.message.reply_text('⚙️ Администратор: /admin — панель управления.')
    await update.message.reply_text(TEXTS[context.user_data['lang']]['choose'],reply_markup=ReplyKeyboardMarkup([[TEXTS['ru']['lang_ru'],TEXTS['uz']['lang_uz']]],resize_keyboard=True))

async def choose_language(update,context,lang):
    context.user_data['lang']=lang
    if not enabled() and update.effective_user.id not in ADMIN_IDS:
        await update.message.reply_text(TEXTS[lang]['disabled'],reply_markup=menu_kb(lang)); return
    await update.message.reply_text(TEXTS[lang]['welcome']+'\n\n'+TEXTS[lang]['payment'].format(price=PRICE,card=CARD_NUMBER),reply_markup=payment_kb(lang))
    context.user_data['state']='payment'

async def go_menu(update,context):
    lang=context.user_data.get('lang','ru')
    # Keep the approved order and current draft in memory. Menu is a pause, not a reset.
    context.user_data['state']='menu'
    await update.message.reply_text(TEXTS[lang]['main'],reply_markup=menu_kb(lang))

async def start_button(update,context):
    uid=update.effective_user.id; lang=context.user_data.get('lang','ru'); order=get_order(uid)
    state=context.user_data.get('state')
    if order and order['status']=='approved':
        context.user_data['lang']=order['lang']; context.user_data['order_id']=order['id']
        if state in STEPS or state in REL_STEPS:
            await update.message.reply_text(tr(context,state, i=context.user_data.get('rel_index',0)+1),reply_markup=form_kb(context.user_data['lang'],optional=state in OPTIONAL,photo=state=='photo'))
            return
        context.user_data['state']='name'; await update.message.reply_text(tr(context,'approved'),reply_markup=form_kb(context.user_data['lang'])); return
    if order and order['status']=='pending':
        await update.message.reply_text(tr(context,'checking'),reply_markup=menu_kb(lang)); return
    if not enabled() and uid not in ADMIN_IDS:
        await update.message.reply_text(tr(context,'disabled'),reply_markup=menu_kb(lang)); return
    await update.message.reply_text(TEXTS[lang]['payment'].format(price=PRICE,card=CARD_NUMBER),reply_markup=payment_kb(lang)); context.user_data['state']='payment'

async def new_button(update,context):
    lang=context.user_data.get('lang','ru'); context.user_data.clear(); context.user_data['lang']=lang
    if not enabled() and update.effective_user.id not in ADMIN_IDS:
        await update.message.reply_text(TEXTS[lang]['disabled'],reply_markup=menu_kb(lang)); return
    await update.message.reply_text(TEXTS[lang]['payment'].format(price=PRICE,card=CARD_NUMBER),reply_markup=payment_kb(lang)); context.user_data['state']='payment'

async def language_button(update,context):
    await update.message.reply_text(TEXTS['ru']['choose'],reply_markup=ReplyKeyboardMarkup([[TEXTS['ru']['lang_ru'],TEXTS['uz']['lang_uz']]],resize_keyboard=True)); context.user_data['state']='choose_language'

async def paid_button(update,context):
    uid=update.effective_user.id; lang=context.user_data.get('lang','ru'); order=get_order(uid)
    if order and order['status']=='pending': await update.message.reply_text(TEXTS[lang]['checking'],reply_markup=menu_kb(lang)); return
    if order and order['status']=='approved':
        context.user_data['order_id']=order['id']; context.user_data['state']='name'; await update.message.reply_text(TEXTS[lang]['approved'],reply_markup=form_kb(lang)); return
    oid=create_order(update.effective_user,lang); context.user_data['order_id']=oid; context.user_data['state']='receipt'
    await update.message.reply_text(TEXTS[lang]['receipt'],reply_markup=form_kb(lang))

async def receipt(update,context):
    oid=context.user_data.get('order_id')
    if not oid: oid=create_order(update.effective_user,context.user_data.get('lang','ru')); context.user_data['order_id']=oid
    cap=f"💳 Новый платёж\nЗаказ #{oid}\nПользователь: {update.effective_user.full_name}\nID: {update.effective_user.id}\nUsername: @{update.effective_user.username or '-'}\nСумма: {PRICE}"
    buttons=InlineKeyboardMarkup([[InlineKeyboardButton('✅ Подтвердить',callback_data=f'approve:{oid}'),InlineKeyboardButton('❌ Отклонить',callback_data=f'reject:{oid}')]])
    for aid in ADMIN_IDS:
        try:
            if update.message.photo: await context.bot.send_photo(aid,update.message.photo[-1].file_id,caption=cap,reply_markup=buttons)
            elif update.message.document: await context.bot.send_document(aid,update.message.document.file_id,caption=cap,reply_markup=buttons)
            else: await context.bot.send_message(aid,cap+'\n⚠️ Пользователь отправил не файл/фото.',reply_markup=buttons)
        except Exception as e: print('admin notification:',e)
    update_order(oid,'pending'); context.user_data['state']='waiting_payment'
    await update.message.reply_text(TEXTS[context.user_data.get('lang','ru')]['receipt_ok'],reply_markup=menu_kb(context.user_data.get('lang','ru')))

async def admin_callback(update,context):
    q=update.callback_query; await q.answer()
    if q.from_user.id not in ADMIN_IDS: return
    action,oid_s=q.data.split(':',1); oid=int(oid_s); c=db(); order=c.execute('SELECT * FROM orders WHERE id=?',(oid,)).fetchone(); c.close()
    if not order: return
    if action=='approve':
        update_order(oid,'approved'); lang=order['lang']; await context.bot.send_message(order['user_id'],TEXTS[lang]['approved'],reply_markup=form_kb(lang)); await q.edit_message_reply_markup(reply_markup=None); await q.message.reply_text(TEXTS['ru']['approved_admin'])
    else:
        update_order(oid,'rejected'); await context.bot.send_message(order['user_id'],TEXTS[order['lang']]['rejected'],reply_markup=menu_kb(order['lang'])); await q.edit_message_reply_markup(reply_markup=None); await q.message.reply_text(TEXTS['ru']['rejected_admin'])

async def admin(update,context):
    if update.effective_user.id not in ADMIN_IDS: await update.message.reply_text(TEXTS['ru']['admin_only']); return
    c=db(); pending=c.execute("SELECT COUNT(*) n FROM orders WHERE status='pending'").fetchone()['n']; c.close(); status=TEXTS['ru']['on_status'] if enabled() else TEXTS['ru']['off_status']
    kb=InlineKeyboardMarkup([[InlineKeyboardButton(TEXTS['ru']['on'],callback_data='bot:on'),InlineKeyboardButton(TEXTS['ru']['off'],callback_data='bot:off')]])
    await update.message.reply_text(TEXTS['ru']['admin_panel'].format(status=status,pending=pending),reply_markup=kb)

async def bot_control(update,context):
    q=update.callback_query; await q.answer()
    if q.from_user.id not in ADMIN_IDS:return
    set_setting('enabled','1' if q.data=='bot:on' else '0'); status=TEXTS['ru']['on_status'] if enabled() else TEXTS['ru']['off_status']; c=db(); n=c.execute("SELECT COUNT(*) n FROM orders WHERE status='pending'").fetchone()['n']; c.close(); await q.edit_message_text(TEXTS['ru']['admin_panel'].format(status=status,pending=n))

async def finish(update,context):
    try:
        data=context.user_data; raw=build_document(data,data.get('photo')); safe=re.sub(r'[\\/:*?"<>|]','_',data['name']); filename=f'Obyektivka - {safe}.docx'; await update.message.reply_document(io.BytesIO(raw),filename=filename,caption=tr(context,'done'),reply_markup=menu_kb(data.get('lang','ru')))
        oid=data.get('order_id');
        if oid:update_order(oid,'completed')
        lang=data.get('lang','ru'); context.user_data.clear(); context.user_data['lang']=lang; context.user_data['state']='menu'
    except Exception as e: await update.message.reply_text(tr(context,'error',e=e),reply_markup=menu_kb(context.user_data.get('lang','ru')))

async def form_message(update,context):
    state=context.user_data.get('state'); lang=context.user_data.get('lang','ru'); text=(update.message.text or '').strip(); norm=normalize(text)
    if state=='receipt': return await receipt(update,context)
    if state=='waiting_payment':
        order=get_order(update.effective_user.id)
        if order and order['status']=='approved': context.user_data['state']='name'; context.user_data['order_id']=order['id']; return await update.message.reply_text(TEXTS[order['lang']]['approved'],reply_markup=form_kb(order['lang']))
        await update.message.reply_text(TEXTS[lang]['checking'],reply_markup=menu_kb(lang)); return
    if state=='payment': return await paid_button(update,context) if norm==normalize(TEXTS[lang]['paid']) else await update.message.reply_text(TEXTS[lang]['payment'].format(price=PRICE,card=CARD_NUMBER),reply_markup=payment_kb(lang))
    if state=='choose_language': return
    if state=='menu' or state is None:
        await update.message.reply_text(TEXTS[lang]['main'],reply_markup=menu_kb(lang)); return
    if state in OPTIONAL and norm in {normalize(TEXTS[lang]['no']),normalize(TEXTS[lang]['skip'])}:
        context.user_data[state]='yo‘q' if lang=='uz' else 'нет'; idx=STEPS.index(state); nxt=STEPS[idx+1]; context.user_data['state']=nxt; await update.message.reply_text(tr(context,nxt),reply_markup=form_kb(lang,optional=nxt in OPTIONAL,photo=nxt=='photo')); return
    if state=='photo':
        if update.message.photo:
            f=await update.message.photo[-1].get_file(); context.user_data['photo']=bytes(await f.download_as_bytearray()); context.user_data['state']='rel_count'; await update.message.reply_text(tr(context,'rel_count'),reply_markup=form_kb(lang)); return
        if norm in {normalize(TEXTS[lang]['skip']),normalize(TEXTS[lang]['no']),"o'tkazib yuborish","yo'q","пропустить","нет"}:
            context.user_data['photo']=None; context.user_data['state']='rel_count'; await update.message.reply_text(tr(context,'rel_count'),reply_markup=form_kb(lang)); return
        await update.message.reply_text(tr(context,'photo'),reply_markup=form_kb(lang,photo=True)); return
    if state=='rel_count':
        try:n=int(norm)
        except: n=-1
        if not 0<=n<=30: await update.message.reply_text(tr(context,'rel_count_error'),reply_markup=form_kb(lang)); return
        context.user_data['relatives']=[]; context.user_data['rel_total']=n; context.user_data['rel_index']=0
        if n==0:return await finish(update,context)
        context.user_data['state']='relation'; await update.message.reply_text(tr(context,'relation',i=1),reply_markup=form_kb(lang)); return
    if state in REL_STEPS:
        key={'relation':'relation','rel_name':'name','rel_birth':'birth','rel_work':'work','rel_address':'address'}[state]; context.user_data.setdefault('rel_tmp',{})[key]=text
        i=REL_STEPS.index(state)
        if i<len(REL_STEPS)-1:
            nxt=REL_STEPS[i+1]; context.user_data['state']=nxt; await update.message.reply_text(tr(context,nxt),reply_markup=form_kb(lang)); return
        context.user_data['relatives'].append(context.user_data.pop('rel_tmp')); context.user_data['rel_index']+=1
        if context.user_data['rel_index']>=context.user_data['rel_total']:return await finish(update,context)
        context.user_data['state']='relation'; await update.message.reply_text(tr(context,'relation',i=context.user_data['rel_index']+1),reply_markup=form_kb(lang)); return
    if state in STEPS:
        context.user_data[state]=text; i=STEPS.index(state)
        if i==len(STEPS)-1: return
        nxt=STEPS[i+1]; context.user_data['state']=nxt; await update.message.reply_text(tr(context,nxt),reply_markup=form_kb(lang,optional=nxt in OPTIONAL,photo=nxt=='photo')); return
    await update.message.reply_text(TEXTS[lang]['main'],reply_markup=menu_kb(lang))

async def router(update,context):
    if not update.message:return
    raw=(update.message.text or '').strip(); n=normalize(raw); lang=context.user_data.get('lang','ru'); L=labels(context)
    # Main navigation is checked BEFORE every form state.
    if n==L['menu']: return await go_menu(update,context)
    if n==L['start']: return await start_button(update,context)
    if n==L['new']: return await new_button(update,context)
    if n==L['language']: return await language_button(update,context)
    if raw==TEXTS['ru']['lang_ru']: return await choose_language(update,context,'ru')
    if raw==TEXTS['uz']['lang_uz']: return await choose_language(update,context,'uz')
    if n==L['paid']: return await paid_button(update,context)
    return await form_message(update,context)

async def myid(update,context): await update.message.reply_text(f'Ваш Telegram ID: {update.effective_user.id}')

async def cancel(update,context):
    lang=context.user_data.get('lang','ru'); context.user_data['state']='menu'; await update.message.reply_text(TEXTS[lang]['main'],reply_markup=menu_kb(lang))

def main():
    if not BOT_TOKEN: raise RuntimeError('Не задан BOT_TOKEN. Укажите его в .env')
    if not TEMPLATE.exists(): raise FileNotFoundError(f'Не найден шаблон: {TEMPLATE}')
    db(); app=Application.builder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler('start',start_cmd)); app.add_handler(CommandHandler('admin',admin)); app.add_handler(CommandHandler('myid',myid)); app.add_handler(CommandHandler('cancel',cancel))
    app.add_handler(CallbackQueryHandler(admin_callback,pattern=r'^(approve|reject):')); app.add_handler(CallbackQueryHandler(bot_control,pattern=r'^bot:'))
    app.add_handler(MessageHandler(filters.ALL & ~filters.COMMAND,router))
    print('BOT STARTED — stable menu/payment/filling version')
    app.run_polling(allowed_updates=Update.ALL_TYPES)

if __name__=='__main__': main()
